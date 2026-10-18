"""Build docs/MVP/balance/territory-model.xlsx (and the 07b map figure).

The territory model is spatial, so this workbook holds its outputs (static
values), not live formulas: parameters, per-season series, every sweep from
territory_sweeps.py, and colored map snapshots. Edit territory_sim.py and
re-run this script to change anything.

    python build_territory_xlsx.py                    # run everything (~12 min)
    python build_territory_xlsx.py --sweeps out.json  # reuse cached sweep results
    python build_territory_xlsx.py --png              # also write territory-maps.png (needs matplotlib)

Needs openpyxl and numpy.
"""
import json
import statistics as st
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent))
import territory_sim as T  # noqa: E402
import territory_sweeps as S  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "territory-model.xlsx"
PNG = Path(__file__).resolve().parents[1] / "territory-maps.png"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="44546A")
BIOME_RGB = {T.OCEAN: (191, 216, 242), T.ICE: (255, 255, 255), T.TUNDRA: (216, 224, 216),
             T.DESERT: (237, 224, 178), T.PEAK: (115, 115, 115), T.MOUNTAIN: (166, 158, 148),
             T.OPEN: (199, 219, 173)}
CIV_RGB = [(230, 159, 0), (86, 180, 233), (0, 158, 115), (204, 121, 167)]  # Okabe–Ito (08 §11)
CIV_NAMES = ["Civ A", "Civ B", "Civ C", "Civ D"]


def hexrgb(c):
    return "%02X%02X%02X" % tuple(int(x) for x in c)


def tile_rgb(m, owner, i):
    b = BIOME_RGB[m.biome[i]]
    o = owner[i]
    if o < 0:
        return b
    c = CIV_RGB[o % len(CIV_RGB)]
    return tuple(0.4 * b[k] + 0.6 * c[k] for k in range(3))


def write_table(ws, rows, r0=1, c0=1, title=None, fmt=None):
    if title:
        ws.cell(r0, c0, title).font = Font(bold=True, size=12)
        r0 += 1
    if not rows:
        return r0
    keys = []
    for r in rows:
        for k in r:
            if k not in keys:
                keys.append(k)
    for j, k in enumerate(keys):
        c = ws.cell(r0, c0 + j, k)
        c.font, c.fill = HEAD, HEADFILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[L(c0 + j)].width = max(ws.column_dimensions[L(c0 + j)].width or 0, 11 if j else 34)
    for i, r in enumerate(rows):
        for j, k in enumerate(keys):
            v = r.get(k)
            if isinstance(v, (list, tuple)):
                v = ", ".join(str(x) for x in v)
            c = ws.cell(r0 + 1 + i, c0 + j, v)
            if isinstance(v, float):
                c.number_format = (fmt or {}).get(k, "0.0%" if (k.startswith(("share", "claimed", "contested", "max_share", "min_share")) and abs(v) <= 1) else "0.0")
    return r0 + len(rows) + 2


def era_table(res):
    ee = {}
    for c in range(len(res["starts"])):
        for e, r in T.era_end_rows(res, c).items():
            ee.setdefault(e, []).append(r)
    rows = []
    for e in sorted(ee):
        r0 = ee[e][0]
        rows.append(dict(era=f"E{e}", end_minute=r0["minute"], pop_per_civ=st.mean(r0["pops"]),
                         settlements=r0["S"][0], owned_tiles_per_civ=st.mean(r0["owned"]),
                         share_min=min(r0["share"]), share_max=max(r0["share"]),
                         claimed=r0["claimed"], contested=r0["contested"]))
    return rows


def contact_rows(res):
    rows = []
    for (a, b), v in sorted(res["contact"].items()):
        rows.append(dict(pair=f"{CIV_NAMES[a]}–{CIV_NAMES[b]}", aware=v.get("aware"), contacted=v.get("contact"),
                         era=S.era_at(res, v.get("contact")), border_touch=v.get("border")))
    return rows


def series_sheet(wb, name, res):
    ws = wb.create_sheet(name)
    n = len(res["starts"])
    hdr = ["minute", "era (Civ A)", "pop/civ"] + [f"S {CIV_NAMES[c]}" for c in range(n)] + \
          [f"owned {CIV_NAMES[c]}" for c in range(n)] + [f"share {CIV_NAMES[c]}" for c in range(n)] + \
          ["claimed", "contested", "civ→civ flips"]
    for j, h in enumerate(hdr):
        c = ws.cell(1, j + 1, h)
        c.font, c.fill = HEAD, HEADFILL
    for i, r in enumerate(res["rows"]):
        vals = [r["minute"], r["eras"][0], st.mean(r["pops"])] + r["S"] + r["owned"] + r["share"] + \
               [r["claimed"], r["contested"], r["flips"]]
        for j, v in enumerate(vals):
            c = ws.cell(i + 2, j + 1, float(v) if not isinstance(v, int) else v)
            if hdr[j].startswith(("share", "claimed", "contested")):
                c.number_format = "0.0%"
    ws.freeze_panes = "B2"
    ch = ScatterChart()
    ch.title = f"Dominion share and claimed land ({name})"
    ch.x_axis.title, ch.y_axis.title = "minute", "share of habitable land"
    ch.height, ch.width = 9, 20
    nrow = len(res["rows"]) + 1
    xs = Reference(ws, min_col=1, min_row=2, max_row=nrow)
    for j, h in enumerate(hdr):
        if h.startswith("share") or h == "claimed":
            ser = Series(Reference(ws, min_col=j + 1, min_row=1, max_row=nrow), xs, title_from_data=True)
            ser.marker.symbol = "none"
            ch.series.append(ser)
    ws.add_chart(ch, L(len(hdr) + 2) + "2")


def map_sheet(wb, name, res, step):
    m = res["map"]
    owner, setts, eras = res["snaps"][step]
    ws = wb.create_sheet(name)
    ws.sheet_view.zoomScale = 50
    for x in range(m.W + 1):
        ws.column_dimensions[L(x + 1)].width = 1.6
    for y in range(m.H):
        ws.row_dimensions[y + 2].height = 9
    minute = step * res["rows"][0]["minute"]
    ws.cell(1, 1, f"{name}: minute {minute:.1f}, eras {['E%d' % e for e in eras]}. Tint = owner (Okabe–Ito), "
                  "● = settlement, ★ = start. Odd rows are offset half a hex in the game (not drawn).").font = Font(bold=True)
    sett = {s: c for c, lst in enumerate(setts) for s in lst}
    for i in range(m.N):
        cell = ws.cell(m.y[i] + 2, m.x[i] + 1)
        cell.fill = PatternFill("solid", fgColor=hexrgb(tile_rgb(m, owner, i)))
        if i in sett:
            cell.value = "★" if i in res["starts"] else "●"
            cell.font = Font(size=6, bold=True)


def params_sheet(wb):
    ws = wb.create_sheet("Params")
    rows = [dict(name=k, value=v[0], note=v[1]) for k, v in T.TP.items()]
    r = write_table(ws, rows, title="Scalars (territory_sim.TP)")
    erows = []
    for k, v in T.TERA.items():
        erows.append(dict(table=k, **{f"E{e}": v[e] for e in range(10)}))
    erows.append(dict(table="Radius (recommended, 07b)", **{f"E{e}": S.REC_RADIUS[e] for e in range(10)}))
    for k, v in S.COUPLED.items():
        erows.append(dict(table=f"balance_sim {k} (coupled, 07b R6)", **{f"E{e}": v[e] for e in range(10)}))
    write_table(ws, erows, r0=r, title="Per-era tables")
    ws.column_dimensions["C"].width = 70


def readme_sheet(wb, base):
    ws = wb.active
    ws.title = "README"
    m = base["paired"]["map"]
    lines = [
        "CivSim — four-civilization territory model (WS-A part 2). Results: docs/MVP/07b-territory-findings.md",
        "",
        "What it is: a deterministic spatial model on a 160×100 hex map (Continents, 40% land, point-symmetric seed 1).",
        "Four civilizations follow the live population and settlement trajectories of balance_sim.py (WS-H demography + R6, see 07c).",
        "Each season the model places fission sites (new land × distance fit, never on rival land), radiates claims",
        "(Emit = sqrt(pop per settlement) × CohesionFactor × ClaimMult, linear falloff to 0 at radius + tail),",
        "resolves ownership (claim >= T and >= (1 + lead) × runner-up, held a full season), and records Dominion",
        "share, contact (figures within reach + reach + 2 tiles by walking distance) and border touches.",
        "",
        f"Map: {m.W}×{m.H} = {m.N} tiles, {int(m.land.sum())} land, habitable weight {m.total_hab:.0f}.",
        "",
        "This workbook holds static outputs (no live formulas): the model is spatial and lives in",
        "tools/territory_sim.py; the scenario sweeps are tools/territory_sweeps.py.",
        "Regenerate: python docs/MVP/balance/tools/build_territory_xlsx.py [--png]",
        "Input version (rebuilt 2026-10-18): the live balance_sim.py/params.py (WS-H + R6), exactly point-symmetric maps",
        "and mirror-invariant tie-breaks (07b §9.5). 07b §3–§7 quote the original WS-A part-1 run; §9 has the current numbers.",
        "Recommended = paired starts 25, T 0.7, radii 3·3·4·4·5·5·5·7·9·9, seasonal camps (07b §9.4).",
        "",
        "Sheets: Params · Eras (era-end tables) · Series_paired / Series_maxspread (per season, with chart) ·",
        "Contact · Early · Ownership · Dominion · MapSize · Coupling · Recommended · Live (07b §9) · Map_* (colored map snapshots)",
    ]
    for i, t in enumerate(lines):
        ws.cell(i + 1, 1, t)
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120


def write_png(rec, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    res = rec[0]
    m = res["map"]
    steps = sorted(res["snaps"])
    fig, axes = plt.subplots(1, len(steps), figsize=(6.2 * len(steps), 4.4))
    for ax, step in zip(axes, steps):
        owner, setts, eras = res["snaps"][step]
        col = np.array([tile_rgb(m, owner, i) for i in range(m.N)]) / 255
        ax.scatter(m.px, -m.py, c=col, marker="h", s=9.5, linewidths=0)
        for c, lst in enumerate(setts):
            ax.scatter(m.px[lst], -m.py[lst], c=[np.array(CIV_RGB[c]) / 255], edgecolors="k", s=10, linewidths=0.5)
        ax.scatter(m.px[res["starts"]], -m.py[res["starts"]], c="k", marker="*", s=40)
        minute = step * 15 / 60
        r = res["rows"][step - 1]
        ax.set_title(f"minute {minute:.0f} (E{eras[0]}): claimed {100 * r['claimed']:.0f}%, "
                     f"shares {', '.join('%.0f' % (100 * x) for x in r['share'])}%", fontsize=9)
        ax.set_aspect("equal")
        ax.axis("off")
    fig.tight_layout()
    fig.savefig(path, dpi=110)
    plt.close(fig)


def main():
    args = sys.argv[1:]
    sweeps = None
    if "--sweeps" in args:
        sweeps = json.loads(Path(args[args.index("--sweeps") + 1]).read_text())
    else:
        sweeps = S.run_all()
    base = S.base_runs()
    rec = S.recommended()
    wb = Workbook()
    readme_sheet(wb, base)
    params_sheet(wb)
    ws = wb.create_sheet("Eras")
    r = write_table(ws, era_table(base["maxspread"]), title="v0.1 starts (max spread, nearest rival 68 tiles), 4× Balanced")
    r = write_table(ws, contact_rows(base["maxspread"]), r0=r, title="Contact, v0.1 starts")
    r = write_table(ws, era_table(base["paired"]), r0=r, title="Paired starts 25 tiles apart (R1), 4× Balanced")
    r = write_table(ws, contact_rows(base["paired"]), r0=r, title="Contact, paired starts")
    r = write_table(ws, S.legacy_by_era(base["paired"]), r0=r, title="Where Dominion Legacy is earned (paired starts)",
                    fmt={"share_of_dominion_legacy": "0.0%"})
    series_sheet(wb, "Series_paired", base["paired"])
    series_sheet(wb, "Series_maxspread", base["maxspread"])
    for g, title in [("contact", "First contact vs. start spacing (median of 4 map seeds)"),
                     ("early", "How full the map is early (median of 4 seeds, paired starts)"),
                     ("ownership", "Ownership constants (seed 1, paired starts)"),
                     ("dominion", "Can anyone reach Dominion 50%?"),
                     ("mapsize", "Map size per player count (~4,000 tiles per civ)"),
                     ("coupling", "balance_sim: live (R6) vs. the v0.1 territory assumption")]:
        ws = wb.create_sheet(g.capitalize() if g != "mapsize" else "MapSize")
        write_table(ws, sweeps[g], title=title)
    ws = wb.create_sheet("Recommended")
    r = 1
    for i, res in enumerate(rec):
        r = write_table(ws, era_table(res), r0=r, title=f"Recommended constants incl. seasonal camps, live trajectory, map seed {S.SEEDS[i]}")
        r = write_table(ws, contact_rows(res), r0=r)
    ws = wb.create_sheet("Live")
    r = 1
    for g, label in [("live", "no seasonal camps"), ("live_camps", "with seasonal camps (adopted)")]:
        for part in ("eras", "seeds", "dominion"):
            r = write_table(ws, sweeps[g][part], r0=r, title=f"07b §9: recommended rules, live trajectory, {label}: {part}")
    for step, label in [(117, "Map_E2end"), (372, "Map_E5end"), (686, "Map_E9")]:
        if step in base["paired"]["snaps"]:
            map_sheet(wb, label, base["paired"], step)
    wb.save(OUT)
    print("wrote", OUT)
    if "--png" in args:
        write_png(rec, PNG)
        print("wrote", PNG)


if __name__ == "__main__":
    main()
