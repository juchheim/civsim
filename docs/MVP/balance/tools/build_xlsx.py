"""Build docs/MVP/balance/balance-model.xlsx with live formulas.

Every Sim and Derived cell is an Excel formula over the Params, Catalog and
Learning sheets, so editing an input re-runs the whole E0-E9 arc in Excel or
Numbers. The formulas mirror balance_sim.py; `--verify` evaluates the workbook
with the `formulas` package and compares it with the Python model.

    python build_xlsx.py            # write the workbook
    python build_xlsx.py --verify   # write, evaluate, compare
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import balance_sim as B  # noqa: E402
import params as P  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "balance-model.xlsx"
SCEN = {"v01": dict(prefix="v01", label="v0.1 (as written)", sim="Sim_v01", der="Derived_v01", rows=2400),
        "tuned": dict(prefix="tun", label="Tuned (recommended)", sim="Sim_tuned", der="Derived_tuned", rows=1200)}

# ---- styles -----------------------------------------------------------------------
INPUT = Font(color="1F4FBF")                 # blue: edit me
DATA = Font(color="6B6B6B", italic=True)     # grey: parsed from the catalog docs
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="2F3B4C")
V01_FILL = PatternFill("solid", fgColor="F6E3C8")
TUN_FILL = PatternFill("solid", fgColor="CDE8E2")
SEL_FILL = PatternFill("solid", fgColor="FFF2A8")
OK_FILL = PatternFill("solid", fgColor="C9E7C1")
BAD_FILL = PatternFill("solid", fgColor="F4C7C3")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
WRAP = Alignment(wrap_text=True, vertical="top")


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def header(ws, row, col, titles, fill=HDR_FILL, font=HDR):
    for i, t in enumerate(titles):
        c = ws.cell(row=row, column=col + i, value=t)
        c.fill, c.font, c.alignment = fill, font, Alignment(wrap_text=True, vertical="center")


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---- Params -------------------------------------------------------------------------
def build_params(wb):
    ws = wb.create_sheet("Params")
    ws["A1"], ws["A1"].font = "Params — every input of the balance model", TITLE
    ws["A2"] = ("Blue = input (edit freely). Both Sim sheets re-run on any change. "
                "v0.1 = numbers as written in docs 01/02/05; Tuned = recommendations in 07-balance-findings.md. "
                "ASSUMED = the docs are silent; the value is a modeling assumption for WS-H/WS-B to replace.")
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 45

    # selectors
    ws["A4"], ws["A4"].font = "Selectors", BOLD
    ws["A5"], ws["B5"] = "Pace", "Standard"
    ws["A6"], ws["B6"] = "Archetype", "Balanced"
    for c in ("B5", "B6"):
        ws[c].fill, ws[c].font = SEL_FILL, INPUT
    dv1 = DataValidation(type="list", formula1='"Standard,Quick,Epic"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(P.ARCHETYPES) + '"', allow_blank=False)
    ws.add_data_validation(dv1), ws.add_data_validation(dv2)
    dv1.add("B5"), dv2.add("B6")
    add_name(wb, "SelPace", "Params!$B$5")
    add_name(wb, "SelArch", "Params!$B$6")
    ws["C5"] = "Quick / Standard / Epic. v0.1 scales research costs (01 §6); Tuned scales the calendar (07 §F7)."
    ws["C6"] = "Sensitivity multipliers applied on top of both scenarios (table below)."

    # pace table
    r0 = 8
    ws.cell(row=r0, column=1, value="Pace table").font = BOLD
    header(ws, r0 + 1, 1, ["Pace", "v0.1 cost ×", "v0.1 s/season", "Tuned cost ×", "Tuned s/season"])
    for i, (pace, d) in enumerate(P.PACES.items()):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=pace)
        for j, v in enumerate([d["v01"][0], d["v01"][1], d["tuned"][0], d["tuned"][1]]):
            ws.cell(row=r, column=2 + j, value=v).font = INPUT
    pr = f"$A${r0 + 2}:$A${r0 + 4}"
    add_name(wb, "v01_PaceCost", "Params!$G$10")
    add_name(wb, "v01_SeasonSeconds", "Params!$G$11")
    add_name(wb, "tun_PaceCost", "Params!$G$12")
    add_name(wb, "tun_SeasonSeconds", "Params!$G$13")
    ws["F9"], ws["F9"].font = "Selected pace →", BOLD
    for k, (lab, col) in enumerate([("v0.1 cost ×", "B"), ("v0.1 s/season", "C"), ("Tuned cost ×", "D"), ("Tuned s/season", "E")]):
        ws.cell(row=10 + k, column=6, value=lab)
        ws.cell(row=10 + k, column=7, value=f"=INDEX(${col}${r0 + 2}:${col}${r0 + 4},MATCH(SelPace,{pr},0))")

    # archetypes
    r0 = 15
    ws.cell(row=r0, column=1, value="Archetype sensitivity multipliers (ASSUMED; 06 §8 asks for all five archetypes)").font = BOLD
    header(ws, r0 + 1, 1, ["Archetype"] + P.ARCH_FIELDS)
    for i, (a, vals) in enumerate(P.ARCHETYPES.items()):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=a)
        for j, v in enumerate(vals):
            ws.cell(row=r, column=2 + j, value=v).font = INPUT
    rsel = r0 + 2 + len(P.ARCHETYPES)
    ws.cell(row=rsel, column=1, value="Selected →").font = BOLD
    names = f"$A${r0 + 2}:$A${rsel - 1}"
    for j, f in enumerate(P.ARCH_FIELDS):
        col = L(2 + j)
        c = ws.cell(row=rsel, column=2 + j, value=f"=INDEX({col}${r0 + 2}:{col}${rsel - 1},MATCH(SelArch,{names},0))")
        c.fill = SEL_FILL
        add_name(wb, f"Arch_{f}", f"Params!${col}${rsel}")

    # scalars
    r0 = rsel + 2
    ws.cell(row=r0, column=1, value="Scalars").font = BOLD
    header(ws, r0 + 1, 1, ["Name", "v0.1", "Tuned", "Unit", "Source / note", "Changed?"])
    for i, (name, (v1, v2, unit, note)) in enumerate(P.SCALARS.items()):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=v1).font = INPUT
        ws.cell(row=r, column=3, value=v2).font = INPUT
        ws.cell(row=r, column=4, value=unit)
        ws.cell(row=r, column=5, value=note)
        ws.cell(row=r, column=6, value=f'=IF(B{r}<>C{r},"changed","")')
        ws.cell(row=r, column=2).fill, ws.cell(row=r, column=3).fill = V01_FILL, TUN_FILL
        if name != "SeasonSeconds":  # season length comes from the pace table
            add_name(wb, f"v01_{name}", f"Params!$B${r}")
            add_name(wb, f"tun_{name}", f"Params!$C${r}")
    r_end = r0 + 2 + len(P.SCALARS)

    # per-era tables
    r0 = r_end + 2
    ws.cell(row=r0, column=1, value="Per-era tables (rows E0–E9). Two columns where v0.1 and Tuned differ.").font = BOLD
    cols = []
    for field, d in P.ERA.items():
        if d["v01"] == d["tuned"]:
            cols.append((field, "both", d["v01"]))
        else:
            cols.append((field, "v01", d["v01"]))
            cols.append((field, "tuned", d["tuned"]))
    header(ws, r0 + 1, 1, ["Era"] + [f"{f}\n{'' if s == 'both' else ('v0.1' if s == 'v01' else 'Tuned')}" for f, s, _ in cols])
    ws.row_dimensions[r0 + 1].height = 32
    for e in range(10):
        ws.cell(row=r0 + 2 + e, column=1, value=P.ERAS[e])
    for j, (field, sc, vals) in enumerate(cols):
        col = 2 + j
        for e, v in enumerate(vals):
            c = ws.cell(row=r0 + 2 + e, column=col, value=v)
            c.font = INPUT
            c.fill = V01_FILL if sc == "v01" else (TUN_FILL if sc == "tuned" else PatternFill())
            if isinstance(v, float) and v < 1:
                c.number_format = "0.0%" if field.startswith("Mort") or field.endswith("Share") else "0.00"
            elif isinstance(v, (int, float)) and v >= 1000:
                c.number_format = "#,##0"
        ref = f"Params!${L(col)}${r0 + 2}:${L(col)}${r0 + 11}"
        if sc in ("both", "v01"):
            add_name(wb, f"v01_{field}", ref)
        if sc in ("both", "tuned"):
            add_name(wb, f"tun_{field}", ref)
    notes_r = r0 + 13
    notes = [
        "EraBase, BreadthN: 02 §8.3 / §13. CostMod: INQ-08 Scientific Method (all costs −10%) from E7. Targets: 02 §8.3 pacing table.",
        "RepScale: representation scale (01 §8, 05 §5.2); Tuned follows modeled population so rendered figures stay ≤ 600 (07 §F6).",
        "FoodCap / HousingCap / GroupCap: v0.1 = one value per era (WS-A placeholders). Tuned = the start value in every row; the Caps sheet raises it project by project (09 §5).",
        "Territory: 09 §8.2 (leader's reachable land on the 160×100 map). ClaimRadius: WS-F (unchanged).",
        "FertMult, MortChild/Adult/Elder (annual hazards outside winter) and WinterMult: tuned = 09 §3–4 (life-table anchors and catalog drivers).",
        "Specialist shares: share of adults in each role (Elder, Shaman, Monastery per settlement); 09 §9.",
    ]
    for i, n in enumerate(notes):
        ws.cell(row=notes_r + i, column=1, value=n)
    widths(ws, {"A": 22, "B": 12, "C": 12, "D": 14, "E": 60, "F": 12})
    for j in range(len(cols)):
        ws.column_dimensions[L(2 + j)].width = max(ws.column_dimensions[L(2 + j)].width or 0, 11)
    ws.freeze_panes = "B4"
    return ws


# ---- Catalog ------------------------------------------------------------------------
def build_catalog(wb):
    ws = wb.create_sheet("Catalog")
    ws["A1"], ws["A1"].font = "Catalog — parsed from 03a/03b/03c by tools/catalog_stats.py", TITLE
    ws["A2"] = ("Grey = data parsed from the catalog docs (regenerate, don't hand-edit). Critical path = keystone closure inside the era, "
                "cheapest side of each one-of. Backfill = earlier-era closure projects not already in a previous keystone's closure.")
    ws.merge_cells("A2:N2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 32
    cols = ["Era", "Keystone", "Projects", "D", "M", "S", "R", "I", "W", "Avg TypeMult", "Crit path n", "Crit path TypeMult Σ",
            "Free-pick avg TypeMult", "Closure n (all eras)", "Forced backfills n"] + [f"Backfill TypeMult Σ from E{s}" for s in range(10)]
    header(ws, 4, 1, cols)
    ws.row_dimensions[4].height = 45
    for e, c in enumerate(B.CAT):
        r = 5 + e
        vals = [P.ERAS[e], f"{c['key']} {c['key_name']}", c["count"]] + [c["types"][t] for t in "DMSRIW"] + [
            c["avg_tm"], c["crit_n"], c["crit_units"], c["free_avg_tm"], c["closure_n"], c["backfill_n"]] + [
            round(u, 3) for u in c["backfill_units"]]
        for j, v in enumerate(vals):
            ws.cell(row=r, column=1 + j, value=v).font = DATA
    for nm, col in [("Cat_CritN", "K"), ("Cat_CritUnits", "L"), ("Cat_FreeTM", "M"), ("Cat_BackfillN", "O")]:
        add_name(wb, nm, f"Catalog!${col}$5:${col}$14")
    add_name(wb, "Cat_BFUnits", "Catalog!$P$5:$Y$14")
    ws.cell(row=16, column=1, value="Validator cross-check vs 04 §2 (in-era critical path / closure):").font = BOLD
    doc = {1: (11, 11), 2: (11, 28), 3: (17, 70), 4: (18, 108), 5: (17, 153), 6: (12, 174), 7: (10, 210), 8: (10, 227), 9: (14, 278), 10: (13, 338)}
    header(ws, 17, 1, ["Keystone", "04 §2 in-era", "Parsed in-era", "04 §2 closure", "Parsed closure", "Match?"])
    for i, c in enumerate(B.CAT):
        r = 18 + i
        k = i + 1
        for j, v in enumerate([c["key"], doc[k][0], c["crit_n"], doc[k][1], c["closure_n"]]):
            ws.cell(row=r, column=1 + j, value=v).font = DATA
        ws.cell(row=r, column=6, value=f'=IF(AND(B{r}=C{r},D{r}=E{r}),"yes","differs")')
    ws.cell(row=29, column=1, value="Critical-path and backfill ids per keystone").font = BOLD
    header(ws, 30, 1, ["Keystone", "In-era critical path ids", "Forced backfill ids"])
    for i, c in enumerate(B.CAT):
        ws.cell(row=31 + i, column=1, value=c["key"]).font = DATA
        ws.cell(row=31 + i, column=2, value=", ".join(c["crit_ids"])).font = DATA
        ws.cell(row=31 + i, column=3, value=", ".join(c["backfill_ids"])).font = DATA
    widths(ws, {"A": 16, "B": 60, "C": 90})
    for j in range(3, 26):
        ws.column_dimensions[L(j + 1)].width = 10
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 10


# ---- Learning -----------------------------------------------------------------------
def build_learning(wb):
    ws = wb.create_sheet("Learning")
    ws["A1"], ws["A1"].font = "Learning — Insight modifiers from the catalog, cumulated per era", TITLE
    ws["A2"] = ("Each row is a catalog effect on Insight. Prob = chance an average civ has it (1.0 = on a keystone closure). "
                "A project counts at half strength in its own era (completed mid-era on average), full strength after. "
                "Multiplier(era, kind) = EXP(Σ weight × LN(1 + pct × prob)).")
    ws.merge_cells("A2:P2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 32
    header(ws, 4, 1, ["Id", "Project", "Era", "Kind", "Pct", "Prob", "LN(1+pct×prob)"] + [f"w·ln E{e}" for e in range(10)])
    n = len(P.LEARNING)
    for i, (pid, name, era, kind, pct, prob) in enumerate(P.LEARNING):
        r = 5 + i
        for j, v in enumerate([pid, name, era, kind, pct, prob]):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.font = INPUT if j >= 4 else DATA
        ws.cell(row=r, column=7, value=f"=LN(1+E{r}*F{r})")
        for e in range(10):
            ws.cell(row=r, column=8 + e, value=f"=IF(C{r}<{e},1,IF(C{r}={e},0.5,0))*G{r}")
    last = 4 + n
    r0 = last + 3
    ws.cell(row=r0 - 1, column=1, value="Cumulative multiplier by era and kind (used by both Sim sheets)").font = BOLD
    header(ws, r0, 1, ["Era"] + P.KINDS)
    for e in range(10):
        r = r0 + 1 + e
        ws.cell(row=r, column=1, value=P.ERAS[e])
        wcol = L(8 + e)
        for k, kind in enumerate(P.KINDS):
            c = ws.cell(row=r, column=2 + k, value=f'=EXP(SUMIF($D$5:$D${last},"{kind}",${wcol}$5:${wcol}${last}))')
            c.number_format = "0.000"
    for k, kind in enumerate(P.KINDS):
        col = L(2 + k)
        add_name(wb, f"LM_{kind}", f"Learning!${col}${r0 + 1}:${col}${r0 + 10}")
    widths(ws, {"A": 9, "B": 26, "C": 6, "D": 11, "E": 7, "F": 7, "G": 9})


# ---- Caps ---------------------------------------------------------------------------
def build_caps(wb):
    ws = wb.create_sheet("Caps")
    ws["A1"], ws["A1"].font = "Caps — population caps raised project by project (09 §5)", TITLE
    ws["A2"] = ("Each row multiplies one settlement cap once research reaches Point = Era + At. Era = the era in which the project typically completes "
                "(a forced backfill completes in the era whose keystone forces it); At = share of that era's requirement done by then. Prob < 1 for optional projects: the cap is × Mult^Prob. Food = sustainable food per settlement; "
                "Lean = late-winter food as a share of Food (× LeanBase, capped at 1); Housing = dwelling capacity; Group = social layer (× CohesionFactor). "
                "Used by both Sim sheets; CapsByProject = 0 (v0.1) switches every multiplier off.")
    ws.merge_cells("A2:J2")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 58
    header(ws, 4, 1, ["Id", "Project", "Done in era", "At", "Cap", "Mult", "Prob", "Point", "Prob × LN(Mult)"])
    n = len(P.CAPS)
    for i, (pid, name, era, at, kind, mult, prob) in enumerate(P.CAPS):
        r = 5 + i
        for j, v in enumerate([pid, name, era, at, kind, mult, prob]):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.font = INPUT if j >= 3 else DATA
        ws.cell(row=r, column=6).number_format = "0.000"
        ws.cell(row=r, column=8, value=f"=C{r}+D{r}").number_format = "0.00"
        ws.cell(row=r, column=9, value=f"=G{r}*LN(F{r})").number_format = "0.0000"
    last = 4 + n
    add_name(wb, "Cap_Point", f"Caps!$H$5:$H${last}")
    add_name(wb, "Cap_Kind", f"Caps!$E$5:$E${last}")
    add_name(wb, "Cap_Ln", f"Caps!$I$5:$I${last}")
    widths(ws, {"A": 9, "B": 32, "C": 6, "D": 7, "E": 9, "F": 8, "G": 7, "H": 8, "I": 14})
    ws.freeze_panes = "C5"


def capmult(kind, pos):
    """Excel: cumulative multiplier of one cap kind at research position `pos`."""
    return f'EXP(X_CapsByProject*SUMPRODUCT((Cap_Point<={pos})*(Cap_Kind="{kind}")*Cap_Ln))'


# ---- Derived (per scenario) ---------------------------------------------------------
DER_COLS = [
    # name, header, formula (row r, era e, prefix X), number format
    ("EraIdx", "Era #", "{e}", "0"),
    ("Unit", "Cost unit × (pace · CostMod · diffusion)", "=X_PaceCost*INDEX(X_CostMod,{i})*X_DiffusionMult*Arch_Diffusion", "0.00"),
    ("PickUnits", "TypeMult Σ of N picks (crit + free)", "=INDEX(Cat_CritUnits,{i})+MAX(0,INDEX(X_BreadthN,{i})-INDEX(Cat_CritN,{i}))*INDEX(Cat_FreeTM,{i})", "0.00"),
    ("AvgCost", "Avg in-era project cost", "=[Unit]{r}*INDEX(X_EraBase,{i})*[PickUnits]{r}/INDEX(X_BreadthN,{i})", "#,##0.0"),
    ("KeyCost", "Keystone cost", "=[Unit]{r}*INDEX(X_EraBase,{i})*X_KeystoneMult", "#,##0"),
    ("KeySpend", "Insight spent while keystone runs (all slots)", "=[KeyCost]{r}*(X_SlotPrimary+X_SlotSecondary+X_SlotTertiary)/X_SlotPrimary", "#,##0"),
    ("Backfill", "Forced backfill cost (EraLag)", "=(1-X_BackfillOverlap)*[Unit]{r}*SUM([BF0]{r}:[BF9]{r})", "#,##0"),
    ("Req", "Era requirement (Insight + Know-how)", "=[Unit]{r}*INDEX(X_EraBase,{i})*[PickUnits]{r}+[KeySpend]{r}+[Backfill]{r}", "#,##0"),
    ("Smax", "Max settlements (territory ÷ claim area)", "=MAX(1,INT(INDEX(X_Territory,{i})*Arch_Territory/((3*INDEX(X_ClaimRadius,{i})*(INDEX(X_ClaimRadius,{i})+1)+1)*X_SettlementPacking)))", "0"),
    ("SettCap", "Settlement cap at era start = MIN(food, housing, group×CF)", "=MIN(INDEX(X_FoodCap,{i})*Arch_FoodCap*" + capmult("Food", "{e}")
     + ",INDEX(X_HousingCap,{i})*" + capmult("Housing", "{e}") + ",INDEX(X_GroupCap,{i})*X_CohesionFactor*" + capmult("Group", "{e}") + ")", "#,##0"),
    ("Boon", "Era Boon Insight ×", "=(1+X_EraBoonPct*X_EraBoonTake)^{e}", "0.000"),
    ("Spill", "Extra in-era completions during keystone", "=([KeySpend]{r}-[KeyCost]{r})/[AvgCost]{r}", "0.0"),
    ("BackfillN", "Forced backfills completed", "=INDEX(Cat_BackfillN,{i})*(1-X_BackfillOverlap)", "0.0"),
    ("Done", "In-era completions (N + spill)", "=INDEX(X_BreadthN,{i})+[Spill]{r}", "0.0"),
] + [(f"BF{s}", f"Backfill cost from E{s}", f"=IF({s}<{{e}},INDEX(Cat_BFUnits,{{i}},{s + 1})*INDEX(X_EraBase,{s + 1})*X_EraLag^({{e}}-{s}),0)", "#,##0")
     for s in range(10)]


def subst(tmpl, colmap, r, p=None, prefix=None, e=None):
    out = re.sub(r"\[(\w+)\]", lambda m: colmap[m.group(1)], tmpl)
    out = out.replace("{r}", str(r))
    if p is not None:
        out = out.replace("{p}", str(p))
    if e is not None:
        out = out.replace("{e}", str(e)).replace("{i}", str(e + 1))
    if prefix:
        out = re.sub(r"\bX_", prefix + "_", out)
    return out


def build_derived(wb, sc):
    cfg = SCEN[sc]
    ws = wb.create_sheet(cfg["der"])
    ws["A1"], ws["A1"].font = f"Derived per-era values — {cfg['label']}", TITLE
    ws["A2"] = "All formulas. Req = N picks at EraBase × TypeMult + keystone phase (all three slots) + forced backfills at EraLag."
    colmap = {name: L(1 + j) for j, (name, *_rest) in enumerate(DER_COLS)}
    header(ws, 4, 1, ["Era"] + [h for _n, h, _f, _fmt in DER_COLS[1:]], fill=V01_FILL if sc == "v01" else TUN_FILL, font=BOLD)
    ws.row_dimensions[4].height = 48
    for e in range(10):
        r = 5 + e
        for j, (name, _h, f, fmt) in enumerate(DER_COLS):
            if name == "EraIdx":
                ws.cell(row=r, column=1, value=P.ERAS[e])
                continue
            c = ws.cell(row=r, column=1 + j, value=subst(f, colmap, r, prefix=cfg["prefix"], e=e))
            c.number_format = fmt
    for name, *_ in DER_COLS[1:]:
        if not name.startswith("BF"):
            add_name(wb, f"{cfg['prefix']}_{name}", f"{cfg['der']}!${colmap[name]}$5:${colmap[name]}$14")
    for j in range(len(DER_COLS)):
        ws.column_dimensions[L(1 + j)].width = 14
    ws.column_dimensions["A"].width = 16
    ws.freeze_panes = "B5"


# ---- Sim (per scenario) -------------------------------------------------------------
CAPF = "SAME"  # row 3 uses the same formula as later rows (caps at research position 0)
SIM_COLS = [
    # key, header, unit, formula (rows ≥ 4), initial row-3 value or formula, number format
    ("t", "Season", "#", "=[t]{p}+1", 0, "0"),
    ("min", "Minute (end of season)", "min", "=[t]{r}*X_SeasonSeconds/60", 0, "0.00"),
    ("yr", "Year", "yr", "=[t]{r}/4", 0, "0.00"),
    ("win", "Winter", "flag", "=IF(MOD([t]{r},4)=0,1,0)", 0, "0"),
    ("adv", "Era advance", "flag", "=IF(AND([era]{p}<=9,[prog]{p}>=[req]{p}),1,0)", 0, "0"),
    ("era", "Era (10 = Apex done)", "#", "=MIN(10,[era]{p}+[adv]{r})", 0, "0"),
    ("ei", "Era row", "index", "=MIN([era]{r},9)+1", 1, "0"),
    ("act", "Active", "flag", "=IF([era]{r}<=9,1,0)", 1, "0"),
    ("pos", "Research position (era + share done)", "era", "=MIN(10,[era]{p}+IF([req]{p}>0,MIN(1,[prog]{p}/[req]{p}),1))", 0, "0.00"),
    ("food", "Food cap / settlement", "people", "=INDEX(X_FoodCap,[ei]{r})*Arch_FoodCap*" + capmult("Food", "[pos]{r}"), CAPF, "#,##0.0"),
    ("lean", "Late-winter food cap", "people", "=[food]{r}*MIN(1,X_LeanBase*" + capmult("Lean", "[pos]{r}") + ")", CAPF, "#,##0.0"),
    ("hous", "Housing cap / settlement", "people", "=INDEX(X_HousingCap,[ei]{r})*" + capmult("Housing", "[pos]{r}"), CAPF, "#,##0"),
    ("grp", "Group cap × CF", "people", "=INDEX(X_GroupCap,[ei]{r})*X_CohesionFactor*" + capmult("Group", "[pos]{r}"), CAPF, "#,##0"),
    ("scap", "Settlement cap = MIN(food, housing, group)", "people", "=MIN([food]{r},[hous]{r},[grp]{r})", CAPF, "#,##0.0"),
    ("cap", "Pop cap (settlements × cap)", "people", "=[S]{p}*[scap]{r}", "=[S]{r}*[scap]{r}", "#,##0"),
    ("crowd", "Crowding", "0–1", "=MIN(1,([pop]{p}/([S]{p}*IF(X_GroupInCrowding=1,[scap]{r},MIN([food]{r},[hous]{r}))))^X_CrowdingExp)", 0, "0.00"),
    ("W", "Season hazard ×", "×", "=IF([win]{r}=1,INDEX(X_WinterMult,[ei]{r}),1)", 1, "0.00"),
    ("births", "Births", "people", "=[act]{r}*[A]{p}*X_WomenShare*X_BaseFertility*INDEX(X_FertMult,[ei]{r})*Arch_Fertility*(1-[crowd]{r})", 0, "0.00"),
    ("dC", "Deaths: children", "people", "=[act]{r}*[C]{p}*INDEX(X_MortChild,[ei]{r})/4*[W]{r}", 0, "0.00"),
    ("dA", "Deaths: adults", "people", "=[act]{r}*[A]{p}*INDEX(X_MortAdult,[ei]{r})*Arch_AdultMort/4*[W]{r}", 0, "0.00"),
    ("dE", "Deaths: elders", "people", "=[act]{r}*[E]{p}*INDEX(X_MortElder,[ei]{r})/4*[W]{r}", 0, "0.00"),
    ("over", "Over-cap loss", "people", "=[act]{r}*MAX(0,[pop]{p}-[cap]{r})*X_OvercapLoss", 0, "0.00"),
    ("aCA", "Child → adult", "people", "=[act]{r}*[C]{p}/(X_ChildYears*4)", 0, "0.00"),
    ("aAE", "Adult → elder", "people", "=[act]{r}*[A]{p}/(X_AdultYears*4)", 0, "0.00"),
    ("C", "Children 0–13", "people", "=[C]{p}+[births]{r}-[dC]{r}-[aCA]{r}-[C]{p}*[over]{r}/[pop]{p}", "=X_StartChildren", "#,##0.0"),
    ("A", "Adults 14–44", "people", "=[A]{p}+[aCA]{r}-[aAE]{r}-[dA]{r}-[A]{p}*[over]{r}/[pop]{p}", "=X_StartAdults", "#,##0.0"),
    ("E", "Elders 45+", "people", "=[E]{p}+[aAE]{r}-[dE]{r}-[E]{p}*[over]{r}/[pop]{p}", "=X_StartElders", "#,##0.0"),
    ("pop", "Population", "people", "=[C]{r}+[A]{r}+[E]{r}", "=[C]{r}+[A]{r}+[E]{r}", "#,##0"),
    ("wtr", "Late-winter food pressure", "flag", "=IF(AND(X_WinterTrigger=1,[win]{r}=1,[pop]{p}/[S]{p}>=[lean]{r}),1,0)", 0, "0"),
    ("fis", "Fission requests", "settlements", "=[act]{r}*IF(AND(OR([pop]{p}/[S]{p}>=X_FissionTrigger*[scap]{r},[wtr]{r}=1),[pop]{p}/[S]{p}*X_DaughterShare>=X_MinDaughter),MAX(1,ROUND([S]{p}*X_FissionRate,0)),0)", 0, "0"),
    ("S", "Settlements", "#", "=IF([act]{r}=1,MAX(1,MIN(INDEX(X_Smax,[ei]{r}),[S]{p}+[fis]{r})),[S]{p})", 1, "0"),
    ("newS", "New settlements", "#", "=MAX(0,[S]{r}-[S]{p})", 0, "0"),
    ("fmin", "Fission minute", "min", '=IF([newS]{r}>0,[min]{r},"")', '=""', "0.00"),
    ("ties", "Insight ties ×", "×", "=IF([era]{r}>=1,1,(1+X_LowTiesShare*([S]{r}-1))/[S]{r})", 1, "0.00"),
    ("eff", "Insight-adults (elders ×2)", "people", "=([A]{r}+X_ElderCommons*INDEX(LM_Elder,[ei]{r})*[E]{r})*[ties]{r}", 0, "#,##0.0"),
    ("commons", "Commons", "Insight/yr", "=[act]{r}*X_CommonsCoef*[eff]{r}^X_CommonsExp*INDEX(LM_Commons,[ei]{r})", 0, "#,##0.0"),
    ("roles", "Roles", "Insight/yr", "=[act]{r}*X_RoleScale*Arch_Roles*([S]{r}*(INDEX(X_ElderRolePerSett,[ei]{r})*X_OutElder*INDEX(LM_Elder,[ei]{r})+INDEX(X_ShamanPerSett,[ei]{r})*X_OutShaman)+[A]{r}*(INDEX(X_PriestShare,[ei]{r})*X_OutPriest+INDEX(X_ScribeShare,[ei]{r})*X_OutScribe*INDEX(LM_Scribe,[ei]{r})+INDEX(X_ScholarShare,[ei]{r})*X_OutScholar*INDEX(LM_Scholar,[ei]{r})+INDEX(X_ResearcherShare,[ei]{r})*X_OutResearcher*INDEX(LM_Researcher,[ei]{r})))", 0, "#,##0.0"),
    ("structs", "Structures", "Insight/yr", "=[act]{r}*X_RoleScale*Arch_Roles*[S]{r}*INDEX(X_MonasteryPerSett,[ei]{r})*X_OutMonastery*INDEX(LM_Monastery,[ei]{r})", 0, "#,##0.0"),
    ("ins", "Insight this season", "Insight", "=([commons]{r}+[roles]{r}+[structs]{r})/4*INDEX(LM_All,[ei]{r})*INDEX(X_Boon,[ei]{r})*Arch_Insight", 0, "#,##0.0"),
    ("khg", "Know-how generated", "KH", "=[act]{r}*(X_KHPerAdult*[A]{r}+X_KHOfCommons*[commons]{r}/4*INDEX(LM_All,[ei]{r}))*X_KHUsable*Arch_KH", 0, "#,##0.0"),
    ("kh", "Know-how applied", "KH", "=MIN([pool]{p}+[khg]{r},[ins]{r}*X_KHMaxRatio)", 0, "#,##0.0"),
    ("pool", "Know-how pool", "KH", "=[pool]{p}+[khg]{r}-[kh]{r}", 0, "#,##0.0"),
    ("progress", "Research progress", "Insight", "=[ins]{r}+[kh]{r}", 0, "#,##0.0"),
    ("req", "Era requirement", "Insight", "=IF([act]{r}=1,INDEX(X_Req,[ei]{r}),0)", "=INDEX(X_Req,1)", "#,##0"),
    ("prog", "Era progress", "Insight", "=IF([act]{r}=1,IF([adv]{r}=1,[prog]{p}-[req]{p},[prog]{p})+[progress]{r},0)", 0, "#,##0.0"),
    ("pct", "Era progress %", "%", "=IF([req]{r}>0,MIN(1,[prog]{r}/[req]{r}),1)", 0, "0%"),
    ("fig", "Rendered figures", "figures", "=[pop]{r}/INDEX(X_RepScale,[ei]{r})", "=[pop]{r}", "#,##0"),
    ("inspm", "Insight per minute", "Insight/min", "=[ins]{r}*60/X_SeasonSeconds", 0, "#,##0.0"),
]
SIM_COL = {k: L(1 + j) for j, (k, *_r) in enumerate(SIM_COLS)}


def build_sim(wb, sc):
    cfg = SCEN[sc]
    ws = wb.create_sheet(cfg["sim"])
    fill = V01_FILL if sc == "v01" else TUN_FILL
    for j, (k, h, unit, *_r) in enumerate(SIM_COLS):
        c = ws.cell(row=1, column=1 + j, value=h)
        c.font, c.fill, c.alignment = BOLD, fill, Alignment(wrap_text=True, vertical="center")
        ws.cell(row=2, column=1 + j, value=unit).font = Font(italic=True, color="777777")
        ws.column_dimensions[L(1 + j)].width = 11
    ws.row_dimensions[1].height = 45
    last = 3 + cfg["rows"]
    for r in range(3, last + 1):
        for j, (k, _h, _u, f, init, fmt) in enumerate(SIM_COLS):
            if r == 3 and init == CAPF:
                v = subst(f, SIM_COL, r, prefix=cfg["prefix"])
            elif r == 3:
                v = init if not isinstance(init, str) else subst(init, SIM_COL, r, prefix=cfg["prefix"])
            else:
                v = subst(f, SIM_COL, r, p=r - 1, prefix=cfg["prefix"])
            c = ws.cell(row=r, column=1 + j, value=v)
            c.number_format = fmt
    ws.freeze_panes = "C3"
    return last


# ---- Results ------------------------------------------------------------------------
RES_COLS = ["Era", "Start (min)", "Minutes", "Target lo", "Target hi", "Duration OK?", "In-era done", "Target done lo",
            "Target done hi", "Done OK?", "Forced backfills", "Sec per completion", "Primary-slot project (min)",
            "Pop at end", "Settlements", "Rendered figures", "≤600?", "Insight/min (avg)", "Know-how share", "Roles share", "Finished?"]


def build_results(wb, lasts):
    ws = wb.create_sheet("Results", 1)
    ws["A1"], ws["A1"].font = "Results — era pacing vs. 02 §8.3 targets (live)", TITLE
    ws["A2"] = "=\"Pace: \"&SelPace&\"   ·   Archetype: \"&SelArch&\"   (change on Params)\""
    blocks = {}
    row = 4
    for sc in ("v01", "tuned"):
        cfg = SCEN[sc]
        X, sim, last = cfg["prefix"], cfg["sim"], lasts[sc]
        fill = V01_FILL if sc == "v01" else TUN_FILL
        ws.cell(row=row, column=1, value=cfg["label"]).font = Font(bold=True, size=12)
        header(ws, row + 1, 1, RES_COLS, fill=fill, font=BOLD)
        ws.row_dimensions[row + 1].height = 45
        rng = lambda k: f"{sim}!${SIM_COL[k]}$4:${SIM_COL[k]}${last}"  # noqa: E731
        first = row + 2
        for e in range(10):
            r = first + e
            i = e + 1
            f = {
                "A": P.ERAS[e],
                "B": f'=COUNTIF({rng("era")},"<"&{e})*{X}_SeasonSeconds/60',
                "C": f'=COUNTIF({rng("era")},{e})*{X}_SeasonSeconds/60',
                "D": f"=INDEX({X}_TargetMinLo,{i})",
                "E": f"=INDEX({X}_TargetMinHi,{i})",
                "F": f'=IF(U{r}<>"yes","unfinished",IF(AND(C{r}>=D{r},C{r}<=E{r}),"✓",IF(C{r}<D{r},"too fast","too slow")))',
                "G": f"=INDEX({X}_Done,{i})",
                "H": f"=INDEX({X}_TargetDoneLo,{i})",
                "I": f"=INDEX({X}_TargetDoneHi,{i})",
                "J": f'=IF(AND(G{r}>=H{r},G{r}<=I{r}),"✓","✗")',
                "K": f"=INDEX({X}_BackfillN,{i})",
                "L": f"=IF(C{r}>0,C{r}*60/(G{r}+K{r}),0)",
                "M": f"=IF(C{r}>0,INDEX({X}_AvgCost,{i})/({X}_SlotPrimary/({X}_SlotPrimary+{X}_SlotSecondary+{X}_SlotTertiary)*(SUMIF({rng('era')},{e},{rng('progress')})/C{r})),0)",
                "N": f'=IF(C{r}>0,INDEX({sim}!${SIM_COL["pop"]}$4:${SIM_COL["pop"]}${last},COUNTIF({rng("era")},"<"&{e})+COUNTIF({rng("era")},{e})),0)',
                "O": f'=IF(C{r}>0,INDEX({sim}!${SIM_COL["S"]}$4:${SIM_COL["S"]}${last},COUNTIF({rng("era")},"<"&{e})+COUNTIF({rng("era")},{e})),0)',
                "P": f"=N{r}/INDEX({X}_RepScale,{i})",
                "Q": f'=IF(P{r}<=600,"✓","✗")',
                "R": f"=IF(C{r}>0,SUMIF({rng('era')},{e},{rng('inspm')})*{X}_SeasonSeconds/60/C{r},0)",
                "S": f"=IF(C{r}>0,SUMIF({rng('era')},{e},{rng('kh')})/SUMIF({rng('era')},{e},{rng('progress')}),0)",
                "T": f"=IF(C{r}>0,(SUMIF({rng('era')},{e},{rng('roles')})+SUMIF({rng('era')},{e},{rng('structs')}))/(SUMIF({rng('era')},{e},{rng('commons')})+SUMIF({rng('era')},{e},{rng('roles')})+SUMIF({rng('era')},{e},{rng('structs')})),0)",
                "U": f'=IF(COUNTIF({rng("era")},">"&{e})>0,"yes",IF(C{r}>0,"running","not reached"))',
            }
            fmts = {"B": "0.0", "C": "0.0", "G": "0.0", "K": "0.0", "L": "0", "M": "0.00", "N": "#,##0", "O": "0",
                    "P": "#,##0", "R": "#,##0", "S": "0%", "T": "0%"}
            for col, v in f.items():
                c = ws[f"{col}{r}"]
                c.value = v
                c.number_format = fmts.get(col, "General")
        tr = first + 10
        ws.cell(row=tr, column=1, value="Total / arc").font = BOLD
        ws[f"B{tr}"] = f'=IF(COUNTIF({rng("era")},10)>0,COUNTIF({rng("era")},"<10")*{X}_SeasonSeconds/60,"not finished")'
        ws[f"C{tr}"] = f"=SUM(C{first}:C{tr - 1})"
        ws[f"D{tr}"] = f"=SUM(D{first}:D{tr - 1})"
        ws[f"E{tr}"] = f"=SUM(E{first}:E{tr - 1})"
        ws[f"G{tr}"] = f"=SUM(G{first}:G{tr - 1})"
        ws[f"K{tr}"] = f"=SUM(K{first}:K{tr - 1})"
        for col in "BCDEGK":
            ws[f"{col}{tr}"].number_format = "0.0"
            ws[f"{col}{tr}"].font = BOLD
        ws[f"F{tr}"] = f'=IF(ISNUMBER(B{tr}),"total "&TEXT(B{tr},"0")&" min (target ≈180 Standard)","")'
        blocks[sc] = (first, tr)
        for col in "FJQ":
            rng_cf = f"{col}{first}:{col}{tr - 1}"
            ws.conditional_formatting.add(rng_cf, FormulaRule(formula=[f'{col}{first}="✓"'], fill=OK_FILL))
            ws.conditional_formatting.add(rng_cf, FormulaRule(formula=[f'AND({col}{first}<>"✓",{col}{first}<>"")'], fill=BAD_FILL))
        row = tr + 3

    # acceptance checks (06 §6)
    ws.cell(row=row, column=1, value="Acceptance checks (06 §6 items 4 and 6; 01 §6 match length)").font = Font(bold=True, size=12)
    header(ws, row + 1, 1, ["Check", "Target", "v0.1", "Tuned", "v0.1 OK?", "Tuned OK?"])
    checks = [
        ("First fission (minute)", "6–10", 6, 10, lambda sc: f'=IF(COUNT({SCEN[sc]["sim"]}!${SIM_COL["fmin"]}$4:${SIM_COL["fmin"]}${lasts[sc]})=0,"none",MIN({SCEN[sc]["sim"]}!${SIM_COL["fmin"]}$4:${SIM_COL["fmin"]}${lasts[sc]}))'),
        ("KEY-1 reached (minute)", "7–9 (±25%)", 7 * 0.75, 9 * 1.25, lambda sc: f"=B{blocks[sc][0] + 1}"),
        ("KEY-2 reached (minute)", "16–20 (±25%)", 16 * 0.75, 20 * 1.25, lambda sc: f"=B{blocks[sc][0] + 2}"),
        ("Full arc (minutes)", "Standard ≈180 · Quick ≈90 · Epic ≈300 (±15%)", None, None, lambda sc: f"=B{blocks[sc][1]}"),
    ]
    for k, (name, tgt, lo, hi, fn) in enumerate(checks):
        r = row + 2 + k
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=tgt)
        ws.cell(row=r, column=3, value=fn("v01")).number_format = "0.0"
        ws.cell(row=r, column=4, value=fn("tuned")).number_format = "0.0"
        for j, col in enumerate("CD"):
            if lo is None:
                f = f'=IF(ISNUMBER({col}{r}),IF(ABS({col}{r}/IF(SelPace="Quick",90,IF(SelPace="Epic",300,180))-1)<=0.15,"✓","✗"),"✗")'
            else:
                f = f'=IF(ISNUMBER({col}{r}),IF(AND({col}{r}>={lo},{col}{r}<={hi}),"✓","✗"),"✗")'
            ws.cell(row=r, column=5 + j, value=f)
    ws.conditional_formatting.add(f"E{row + 2}:F{row + 5}", FormulaRule(formula=[f'E{row + 2}="✓"'], fill=OK_FILL))
    ws.conditional_formatting.add(f"E{row + 2}:F{row + 5}", FormulaRule(formula=[f'E{row + 2}="✗"'], fill=BAD_FILL))
    widths(ws, {"A": 26, "B": 12, "F": 14, "J": 9, "Q": 7, "U": 11})
    for col in "CDEGHIKLMNOPRST":
        ws.column_dimensions[col].width = 11
    ws.freeze_panes = "B4"
    return blocks


# ---- Charts -------------------------------------------------------------------------
def build_charts(wb, blocks, lasts):
    ws = wb.create_sheet("Charts", 2)
    ws["A1"], ws["A1"].font = "Charts (live)", TITLE
    res = wb["Results"]
    # helper table: era minutes v01 / tuned / target midpoint
    header(ws, 3, 1, ["Era", "v0.1 minutes", "Tuned minutes", "Target midpoint"])
    for e in range(10):
        r = 4 + e
        ws.cell(row=r, column=1, value=P.ERAS[e].split()[0])
        ws.cell(row=r, column=2, value=f"=Results!C{blocks['v01'][0] + e}")
        ws.cell(row=r, column=3, value=f"=Results!C{blocks['tuned'][0] + e}")
        ws.cell(row=r, column=4, value=f"=(Results!D{blocks['tuned'][0] + e}+Results!E{blocks['tuned'][0] + e})/2")
    bar = BarChart()
    bar.type, bar.grouping = "col", "clustered"
    bar.title, bar.y_axis.title, bar.x_axis.title = "Era duration (minutes) vs target", "minutes", "era"
    bar.add_data(Reference(ws, min_col=2, max_col=4, min_row=3, max_row=13), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=4, max_row=13))
    bar.height, bar.width = 9, 18
    ws.add_chart(bar, "F3")

    def scatter(title, ykey, ytitle, anchor, log=True):
        ch = ScatterChart()
        ch.title, ch.style = title, 13
        ch.x_axis.title, ch.y_axis.title = "minute", ytitle
        for sc in ("v01", "tuned"):
            sim = wb[SCEN[sc]["sim"]]
            xcol = list(SIM_COL).index("min") + 1
            ycol = list(SIM_COL).index(ykey) + 1
            xs = Reference(sim, min_col=xcol, min_row=4, max_row=lasts[sc])
            ys = Reference(sim, min_col=ycol, min_row=4, max_row=lasts[sc])
            s = Series(ys, xs, title=SCEN[sc]["label"])
            s.marker.symbol = "none"
            s.smooth = False
            ch.series.append(s)
        if log:
            ch.y_axis.scaling.logBase = 10
        ch.x_axis.scaling.min = 0
        ch.x_axis.scaling.max = 360
        ch.height, ch.width = 9, 18
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ws.add_chart(ch, anchor)

    scatter("Era reached over time", "era", "era", "F22", log=False)
    scatter("Population", "pop", "people (log)", "F41")
    scatter("Insight per minute", "inspm", "Insight/min (log)", "F60")
    ws.column_dimensions["A"].width = 8
    for c in "BCD":
        ws.column_dimensions[c].width = 13


# ---- Sweeps (static snapshot) -------------------------------------------------------
def build_sweeps(wb):
    ws = wb.create_sheet("Sweeps")
    ws["A1"], ws["A1"].font = "Sweeps — static snapshot from tools/balance_sim.py (not live)", TITLE
    ws["A2"] = ("Every pace × archetype for both scenarios, computed by the Python mirror of these formulas at build time. "
                "The live sheets show one Pace/Archetype at a time (Params selectors); use this table to compare them all.")
    ws.merge_cells("A2:P2")
    ws["A2"].alignment = WRAP
    header(ws, 4, 1, ["Scenario", "Pace", "Archetype", "Total (min)", "Target (min)", "First fission"] + [f"E{e} min" for e in range(10)])
    r = 5
    for sc in ("v01", "tuned"):
        for pace in P.PACES:
            for arch in P.ARCHETYPES:
                s = P.scenario(sc, pace, arch)
                rows, D = B.run(s)
                res, ff, tot = B.summarize(rows, D, s)
                vals = [SCEN[sc]["label"], pace, arch, tot if tot else "not finished", {"Standard": 180, "Quick": 90, "Epic": 300}[pace],
                        ff] + [round(x["minutes"], 2) if x else None for x in res]
                for j, v in enumerate(vals):
                    c = ws.cell(row=r, column=1 + j, value=v)
                    c.font = DATA
                    if j >= 3 and isinstance(v, float):
                        c.number_format = "0.0"
                r += 1
    widths(ws, {"A": 20, "B": 10, "C": 11, "D": 11, "E": 11, "F": 12})


# ---- README -------------------------------------------------------------------------
def build_readme(wb, blocks):
    ws = wb.active
    ws.title = "README"
    ws["A1"], ws["A1"].font = "CivSim balance model — WS-A (design only)", Font(bold=True, size=16)
    lines = [
        ("What", "A season-by-season model of one civilization from E0 Dawn to the KEY-10 Apex: cohorts, births, deaths, the three population caps and fission (05 §9), "
                 "Insight (02 §8.1), the 4:2:1 slots (02 §8.2), project costs from EraBase × TypeMult (02 §8.3), Know-how (02 §9) and keystone breadth N (02 §13, 04 §2)."),
        ("Why", "To check the pacing targets in 02 §8.3 (eras ≈ 8 → 30 min, ≈ 3 h total at Standard, 13–25 completions per era) and tune the constants. Findings: 07-balance-findings.md."),
        ("Two scenarios", "v0.1 = the numbers as written. Tuned = the recommended constants. Both run side by side (Sim_v01, Sim_tuned) from the same Params sheet."),
        ("How to use", "Edit blue cells on Params. Pick Pace and Archetype at the top of Params. Everything else is formulas; Results and Charts update. Open in Excel or Numbers (recalculation on open)."),
        ("Sheets", "Results: pacing vs targets + acceptance checks · Charts · Params: all inputs · Catalog: per-era counts, critical paths, forced backfills parsed from 03a–c · "
                   "Learning: Insight modifiers from the catalog · Caps: population caps by project · Derived_*: per-era costs and requirements · Sim_*: one row per season (15 s) · Sweeps: static archetype × pace table."),
        ("Research model", "Each era requires: N picks (in-era critical path at its exact TypeMult + free picks at the era's average) × EraBase, plus the keystone phase "
                           "(keystone in Primary at 4/7; the other two slots keep completing projects, which is where completions above N come from), plus forced backfills of "
                           "earlier-era closure projects at EraLag 0.7^Δ. Progress = Insight + Know-how; overflow carries into the next era."),
        ("Population model", "Three cohorts (children 0–13, adults 14–44, elders 45+), seasonal hazards with a winter multiplier (09-population-and-health.md). "
                             "Caps per settlement rise project by project as research advances (Caps sheet): food, late-winter food, housing, group × Cohesion. "
                             "Births are damped by crowding against food and housing. A settlement fissions when its average size passes 85% of MIN(food, housing, group), "
                             "or when a late-winter check finds it larger than its lean-season food, and the daughter band would be viable, up to the territory limit. "
                             "Mean-field: one 'average settlement', not a spatial sim."),
        ("Color code", "Blue = input · grey italic = data parsed from the catalog docs · black = formula · orange = v0.1 · teal = tuned · yellow = selector."),
        ("Regenerate", "docs/MVP/balance/tools/build_xlsx.py rebuilds this file from tools/params.py and the catalog docs; --verify re-evaluates it against the Python model."),
        ("Limits", "Single civilization (no diffusion from rivals, no war losses, no trade). Mean-field settlements. Epidemics enter as their expected loss inside the "
                   "per-era hazards (09 §10), not as outbreaks. Effects of individual projects beyond Insight and caps are not modeled."),
    ]
    for i, (k, v) in enumerate(lines):
        ws.cell(row=3 + i, column=1, value=k).font = BOLD
        c = ws.cell(row=3 + i, column=2, value=v)
        c.alignment = WRAP
        ws.row_dimensions[3 + i].height = 48
    r = 4 + len(lines)
    ws.cell(row=r, column=1, value="Headline (live)").font = Font(bold=True, size=12)
    ws.cell(row=r + 1, column=1, value="v0.1 full arc")
    ws.cell(row=r + 1, column=2, value=f'=IF(ISNUMBER(Results!B{blocks["v01"][1]}),TEXT(Results!B{blocks["v01"][1]},"0")&" min","not finished")&"   ·   E0 "&TEXT(Results!C{blocks["v01"][0]},"0.0")&" min, E9 "&TEXT(Results!C{blocks["v01"][0] + 9},"0.0")&" min"')
    ws.cell(row=r + 2, column=1, value="Tuned full arc")
    ws.cell(row=r + 2, column=2, value=f'=IF(ISNUMBER(Results!B{blocks["tuned"][1]}),TEXT(Results!B{blocks["tuned"][1]},"0")&" min","not finished")&"   ·   E0 "&TEXT(Results!C{blocks["tuned"][0]},"0.0")&" min, E9 "&TEXT(Results!C{blocks["tuned"][0] + 9},"0.0")&" min"')
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 150


def build(out=OUT, rows=None):
    if rows:
        for sc in SCEN:
            SCEN[sc]["rows"] = rows
    wb = Workbook()
    build_params(wb)
    build_catalog(wb)
    build_learning(wb)
    build_caps(wb)
    lasts = {}
    for sc in ("v01", "tuned"):
        build_derived(wb, sc)
    for sc in ("v01", "tuned"):
        lasts[sc] = build_sim(wb, sc)
    blocks = build_results(wb, lasts)
    build_charts(wb, blocks, lasts)
    build_sweeps(wb)
    build_readme(wb, blocks)
    order = ["README", "Results", "Charts", "Params", "Catalog", "Learning", "Caps", "Derived_v01", "Derived_tuned", "Sim_v01", "Sim_tuned", "Sweeps"]
    wb._sheets = [wb[n] for n in order]
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("wrote", out)
    return lasts, blocks


def verify(path, lasts, blocks):
    """Evaluate a workbook with `formulas` and compare it with balance_sim."""
    import formulas
    import time
    t0 = time.time()
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    print(f"evaluated in {time.time() - t0:.0f} s")
    index = {k.upper(): v for k, v in sol.items()}

    def val(sheet, cell):
        key = f"'[{path.name}]{sheet}'!{cell}".upper()
        v = index[key].value
        while hasattr(v, "__getitem__") and not isinstance(v, str):
            v = v[0]
        return v

    worst = 0.0
    for sc in ("v01", "tuned"):
        s = P.scenario(sc)
        rows, D = B.run(s, SCEN[sc]["rows"])
        for t in (1, 2, 10, 50, 100, 200, 400, 688, 1000, 1400, 2000):
            if t > SCEN[sc]["rows"]:
                continue
            r = 3 + t
            for key, pk in (("pop", "pop"), ("S", "S"), ("ins", "insight"), ("prog", "prog"), ("era", "era")):
                got = float(val(SCEN[sc]["sim"], f"{SIM_COL[key]}{r}"))
                want = float(rows[t][pk])
                rel = abs(got - want) / max(1e-9, abs(want)) if want else abs(got)
                worst = max(worst, rel)
                if rel > 1e-6:
                    print(f"MISMATCH {sc} t={t} {key}: sheet {got} vs python {want}")
                elif t in (100, 688) and key in ("pop", "ins", "era"):
                    print(f"  ok {sc} t={t} {key}: sheet {got:.6g} = python {want:.6g}")
        res, ff, tot = B.summarize(rows, D, s)
        print(f"{sc}: sheet total {val('Results', 'B' + str(blocks[sc][1]))}  python total {tot}")
        for e in range(10):
            got = val("Results", f"C{blocks[sc][0] + e}")
            want = res[e]["minutes"] if res[e] else 0
            if abs(float(got) - want) > 1e-9:
                print(f"MISMATCH {sc} E{e} minutes: sheet {got} vs python {want}")
        print(f"  {sc} era minutes sheet:", [round(float(val('Results', f'C{blocks[sc][0] + e}')), 2) for e in range(10)])
        print(f"  {sc} first fission sheet:", val("Results", f"C{blocks['tuned'][1] + 5}" if sc == "v01" else f"D{blocks['tuned'][1] + 5}"), "python:", ff)
    print("worst relative difference:", worst)


if __name__ == "__main__":
    if "--verify" in sys.argv:
        # evaluate a short copy (the evaluator is slow on the full 3,600-row workbook)
        n = int(sys.argv[sys.argv.index("--verify") + 1]) if len(sys.argv) > sys.argv.index("--verify") + 1 else 300
        import tempfile
        tmp = Path(sys.argv[-1]) if sys.argv[-1].endswith(".xlsx") else Path(tempfile.gettempdir()) / "civsim_verify.xlsx"
        lasts, blocks = build(tmp, rows=n)
        verify(tmp, lasts, blocks)
    else:
        build()
